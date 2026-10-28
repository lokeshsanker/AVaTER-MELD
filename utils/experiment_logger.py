import os
from datetime import datetime
from openpyxl import Workbook, load_workbook


RESULTS_DIR = "results"
REPORT_FILE = os.path.join(RESULTS_DIR, "experiment_report.xlsx")

EMOTION_LABELS = [
    "neutral",
    "joy",
    "sadness",
    "anger",
    "fear",
    "disgust",
    "surprise"
]


def _get_or_create_workbook():
    os.makedirs(RESULTS_DIR, exist_ok=True)

    if os.path.exists(REPORT_FILE):
        wb = load_workbook(REPORT_FILE)
    else:
        wb = Workbook()

    # ---------------------------------------------------------
    # Run Summary
    # ---------------------------------------------------------
    if "Run Summary" not in wb.sheetnames:
        ws = wb.create_sheet("Run Summary")
        ws.append([
            "Run ID",
            "Date/Time",
            "Modality",
            "Model",
            "Dataset",
            "Train Samples",
            "Validation Samples",
            "Test Samples",
            "Best Epoch",
            "Accuracy",
            "Macro-F1",
            "Weighted-F1",
            "Checkpoint",
            "Status",
            "Notes"
        ])

    # ---------------------------------------------------------
    # Epoch History
    # ---------------------------------------------------------
    if "Epoch History" not in wb.sheetnames:
        ws = wb.create_sheet("Epoch History")
        ws.append([
            "Run ID",
            "Modality",
            "Model",
            "Epoch",
            "Train Loss",
            "Train Accuracy",
            "Val Loss",
            "Val Accuracy",
            "Macro-F1"
        ])

    # ---------------------------------------------------------
    # Class Results
    # ---------------------------------------------------------
    if "Class Results" not in wb.sheetnames:
        ws = wb.create_sheet("Class Results")
        ws.append([
            "Run ID",
            "Modality",
            "Model",
            "Emotion",
            "Precision",
            "Recall",
            "F1",
            "Support"
        ])

    # ---------------------------------------------------------
    # Configuration
    # ---------------------------------------------------------
    if "Configuration" not in wb.sheetnames:
        ws = wb.create_sheet("Configuration")
        ws.append([
            "Run ID",
            "Modality",
            "Model",
            "Parameter",
            "Value"
        ])

    # ---------------------------------------------------------
    # Confusion Matrix
    # ---------------------------------------------------------
    if "Confusion Matrix" not in wb.sheetnames:
        ws = wb.create_sheet("Confusion Matrix")
        ws.append([
            "Run ID",
            "Modality",
            "Model",
            "Actual / Predicted",
            "neutral",
            "joy",
            "sadness",
            "anger",
            "fear",
            "disgust",
            "surprise"
        ])

    # ---------------------------------------------------------
    # NEW: Class Detection
    # ---------------------------------------------------------
    if "Class Detection" not in wb.sheetnames:
        ws = wb.create_sheet("Class Detection")
        ws.append([
            "Run ID",
            "Version",
            "Modality",
            "Model",
            "Emotion",
            "Target Class",
            "Detected",
            "Detection Count"
        ])

    return wb


def _next_run_id(wb):
    ws = wb["Run Summary"]

    if ws.max_row <= 1:
        return "RUN_001"

    return f"RUN_{ws.max_row:03d}"


def log_run(
    modality,
    model,
    dataset="MELD",
    train_samples="",
    validation_samples="",
    test_samples="",
    best_epoch="",
    accuracy="",
    macro_f1="",
    weighted_f1="",
    checkpoint="",
    status="Complete",
    notes=""
):
    """
    Add one completed experiment to the Excel report.

    Returns:
        run_id
    """

    wb = _get_or_create_workbook()
    run_id = _next_run_id(wb)

    ws = wb["Run Summary"]

    ws.append([
        run_id,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        modality,
        model,
        dataset,
        train_samples,
        validation_samples,
        test_samples,
        best_epoch,
        accuracy,
        macro_f1,
        weighted_f1,
        checkpoint,
        status,
        notes
    ])

    wb.save(REPORT_FILE)

    print(f"[Excel] Experiment logged: {run_id}")
    print(f"[Excel] Report: {REPORT_FILE}")

    return run_id


def log_epoch(
    run_id,
    modality,
    model,
    epoch,
    train_loss="",
    train_accuracy="",
    val_loss="",
    val_accuracy="",
    macro_f1=""
):
    """
    Add one epoch result to the Epoch History sheet.
    """

    wb = _get_or_create_workbook()

    ws = wb["Epoch History"]

    ws.append([
        run_id,
        modality,
        model,
        epoch,
        train_loss,
        train_accuracy,
        val_loss,
        val_accuracy,
        macro_f1
    ])

    wb.save(REPORT_FILE)


def log_class_results(
    run_id,
    modality,
    model,
    class_results
):
    """
    class_results should be a list of dictionaries:

    [
        {
            "emotion": "neutral",
            "precision": 0.50,
            "recall": 0.80,
            "f1": 0.62,
            "support": 469
        }
    ]
    """

    wb = _get_or_create_workbook()

    ws = wb["Class Results"]

    for result in class_results:
        ws.append([
            run_id,
            modality,
            model,
            result.get("emotion", ""),
            result.get("precision", ""),
            result.get("recall", ""),
            result.get("f1", ""),
            result.get("support", "")
        ])

    wb.save(REPORT_FILE)


def log_configuration(
    run_id,
    modality,
    model,
    configuration
):
    """
    configuration should be a dictionary:

    {
        "batch_size": 16,
        "learning_rate": 0.0002
    }
    """

    wb = _get_or_create_workbook()

    ws = wb["Configuration"]

    for parameter, value in configuration.items():
        ws.append([
            run_id,
            modality,
            model,
            parameter,
            value
        ])

    wb.save(REPORT_FILE)


def log_confusion_matrix(
    run_id,
    modality,
    model,
    matrix,
    emotion_labels
):
    """
    Save a confusion matrix to the Excel report.
    """

    wb = _get_or_create_workbook()

    ws = wb["Confusion Matrix"]

    for i, emotion in enumerate(emotion_labels):
        row = [
            run_id,
            modality,
            model,
            emotion
        ]

        row.extend(matrix[i])

        ws.append(row)

    wb.save(REPORT_FILE)


def log_class_detection(
    run_id,
    version,
    modality,
    model,
    emotion_labels,
    predicted_counts
):
    """
    Record which emotion classes were actually detected.

    predicted_counts should contain the number of predictions
    made for each emotion.

    Example:

    predicted_counts = {
        "neutral": 759,
        "joy": 349,
        "sadness": 0,
        "anger": 0,
        "fear": 0,
        "disgust": 0,
        "surprise": 0
    }
    """

    wb = _get_or_create_workbook()

    ws = wb["Class Detection"]

    for emotion in emotion_labels:

        count = int(predicted_counts.get(emotion, 0))

        detected = "Yes" if count > 0 else "No"

        ws.append([
            run_id,
            version,
            modality,
            model,
            emotion,
            "Yes",
            detected,
            count
        ])

    wb.save(REPORT_FILE)

    print(f"[Excel] Class detection recorded for {run_id}")


if __name__ == "__main__":
    print("Experiment logger ready.")
    print(f"Report file: {REPORT_FILE}")